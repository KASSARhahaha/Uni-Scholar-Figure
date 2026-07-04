"""Feature 3: Layer Stack diagram.

Stacks N labeled rectangles vertically with arrows pointing down — the classic
"layered architecture" figure (UI → Service → Data → Storage).
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

DEFAULT_PALETTE = [
    "4F86C6",  # blue
    "5DA5DA",  # cyan
    "60BD68",  # green
    "F17CB0",  # pink
    "FAA43A",  # orange
    "B276B2",  # purple
    "DECF3F",  # yellow
]


def render_layer_stack(
    deck_path: Path,
    layers: list[str],
    *,
    title: str | None = None,
    box_height_in: float = 0.8,
    box_width_in: float = 6.5,
    gap_in: float = 0.25,
    arrow_size_in: float = 0.18,
    palette: list[str] | None = None,
) -> Presentation:
    prs = Presentation(str(deck_path)) if deck_path.exists() else Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    if title:
        slide.shapes.title.text = title

    palette = palette or DEFAULT_PALETTE
    left = Inches((10 - box_width_in) / 2)  # centered on 10in wide
    top0 = Inches(1.5)
    bh = Inches(box_height_in)
    bw = Inches(box_width_in)
    gap = Inches(gap_in)

    for i, label in enumerate(layers):
        top = Inches(1.5 + i * (box_height_in + gap_in))
        color_hex = palette[i % len(palette)]
        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, bw, bh)
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor.from_string(color_hex)
        rect.line.color.rgb = RGBColor.from_string("333333")
        tf = rect.text_frame
        tf.text = label
        for para in tf.paragraphs:
            para.alignment = 2  # center
            for run in para.runs:
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string("FFFFFF")

        # arrow below (skip after last layer)
        if i < len(layers) - 1:
            arrow_top = top + bh
            arrow_h = Inches(arrow_size_in)
            arrow_w = Inches(arrow_size_in)
            arrow_left = left + Inches((box_width_in - arrow_size_in) / 2)
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, arrow_left, arrow_top, arrow_w, arrow_h)
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor.from_string("666666")
            arrow.line.fill.background()

    prs.save(str(deck_path))
    return prs
