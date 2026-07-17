"""Feature 4: Matrix row-offset.

Lays out a RxC grid of equal-size shapes where every other row is shifted by
a configurable Δx — produces the classic "brick pattern" used for nested
matrix / attention-visualization figures.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def render_matrix(
    deck_path: Path,
    *,
    rows: int,
    cols: int,
    title: str | None = None,
    cell_in: float = 1.0,
    gap_in: float = 0.12,
    row_offset_in: float = 0.5,
    offset_mode: str = "alternate",  # alternate | progressive | none
    color_hex: str = "4F86C6",
    label_fn=None,
) -> Presentation:
    prs = Presentation(str(deck_path)) if deck_path.exists() else Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    if title:
        slide.shapes.title.text = title

    cw = Inches(cell_in)
    ch = Inches(cell_in)
    stride_x = cell_in + gap_in
    stride_y = cell_in + gap_in

    total_width = cols * stride_x + abs(row_offset_in)
    start_left = (10 - total_width) / 2  # center
    top0 = 1.5

    for r in range(rows):
        if offset_mode == "alternate":
            dx = row_offset_in if r % 2 == 1 else 0.0
        elif offset_mode == "progressive":
            dx = row_offset_in * r
        else:
            dx = 0.0
        for c in range(cols):
            left = Inches(start_left + dx + c * stride_x)
            top = Inches(top0 + r * stride_y)
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, cw, ch)
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(color_hex)
            shape.line.color.rgb = RGBColor.from_string("FFFFFF")
            if label_fn is not None:
                tf = shape.text_frame
                tf.text = str(label_fn(r, c))
                for para in tf.paragraphs:
                    para.alignment = 2
                    for run in para.runs:
                        run.font.size = Pt(11)
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")

    prs.save(str(deck_path))
    return prs
