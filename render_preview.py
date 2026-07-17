"""Render a .pptx to PNG thumbnails using Pillow (no LibreOffice needed).

This is a faithful-enough preview for shapes our CLI creates: rectangles,
rounded rectangles, ovals, down-arrows, tables, and text. It is NOT a general
pptx renderer — only what we need to eyeball our own output.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

# 10in x 7.5in slide → 96 dpi
DPI = 96
SLIDE_W = int(10 * DPI)
SLIDE_H = int(7.5 * DPI)

PALETTE_FALLBACK = "4F86C6"


def _emu_to_px(emu):
    return int(emu / 914400 * DPI)


def _load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ):
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    return ImageFont.load_default()


def _hex_color(rgb) -> tuple[int, int, int]:
    if rgb is None:
        return (0x4F, 0x86, 0xC6)
    s = str(rgb)
    return (int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _draw_shape(draw, shape, img):
    try:
        left = _emu_to_px(shape.left)
        top = _emu_to_px(shape.top)
        w = _emu_to_px(shape.width)
        h = _emu_to_px(shape.height)
    except Exception:
        return
    # Table
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
        tbl = shape.table
        rows = list(tbl.rows)
        cols = list(tbl.columns)
        col_w = [max(1, _emu_to_px(c.width)) for c in cols]
        row_h = [max(1, _emu_to_px(r.height)) for r in rows]
        x = left
        for ci, cw in enumerate(col_w):
            y = top
            for ri, rh in enumerate(row_h):
                try:
                    txt = tbl.cell(ri, ci).text
                except Exception:
                    txt = ""
                # alt-row band
                fill = (240, 240, 245) if ri == 0 else (255, 255, 255) if ri % 2 == 1 else (245, 245, 250)
                draw.rectangle([x, y, x + cw, y + rh], fill=fill, outline=(120, 120, 120))
                if txt:
                    fnt = _load_font(min(14, max(10, rh // 2)))
                    draw.text((x + 4, y + 4), txt[:20], fill=(0, 0, 0), font=fnt)
            y += rh
            x += cw
        return

    # Auto shape (rectangle / rounded rect / oval / arrow)
    try:
        fill = shape.fill
        fill_rgb = None
        if fill.type == 1:  # solid
            fill_rgb = _hex_color(fill.fore_color.rgb)
    except Exception:
        fill_rgb = None
    fill_rgb = fill_rgb or _hex_color(None)

    # Auto shape sub-type (oval / arrow / rounded rect / etc.)
    auto_kind = None
    try:
        auto_kind = shape.auto_shape_type
    except Exception:
        pass

    if auto_kind == MSO_SHAPE.OVAL:
        draw.ellipse([left, top, left + w, top + h], fill=fill_rgb)
    elif auto_kind == MSO_SHAPE.DOWN_ARROW:
        cx = left + w // 2
        draw.polygon(
            [(left, top), (left + w, top), (cx, top + h)],
            fill=fill_rgb,
        )
    elif auto_kind == MSO_SHAPE.RECTANGLE:
        draw.rectangle([left, top, left + w, top + h], fill=fill_rgb)
    else:
        # Rounded rectangle (default for our ROUNDED_RECTANGLE shapes)
        radius = min(15, w // 4, h // 4)
        try:
            draw.rounded_rectangle([left, top, left + w, top + h], radius=radius, fill=fill_rgb)
        except Exception:
            draw.rectangle([left, top, left + w, top + h], fill=fill_rgb)

    # Text
    if shape.has_text_frame:
        txt = shape.text_frame.text
        if txt:
            fnt_size = min(20, max(10, h // 3))
            fnt = _load_font(fnt_size)
            # Center text
            cx = left + w // 2
            cy = top + h // 2
            try:
                bbox = draw.textbbox((0, 0), txt, font=fnt)
                tw = bbox[2] - bbox[0]
                th = bbox[3] - bbox[1]
            except Exception:
                tw, th = fnt_size * len(txt) // 2, fnt_size
            # White text on colored shape
            draw.text((cx - tw // 2, cy - th // 2), txt, fill=(255, 255, 255), font=fnt)


def _draw_title(draw, slide):
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        txt = slide.shapes.title.text_frame.text
        if txt:
            fnt = _load_font(28)
            draw.text((int(0.5 * DPI), int(0.3 * DPI)), txt, fill=(0x1F, 0x4E, 0x79), font=fnt)


def render_slide(slide, idx: int, out_dir: Path) -> Path:
    img = Image.new("RGB", (SLIDE_W, SLIDE_H), "white")
    draw = ImageDraw.Draw(img)
    _draw_title(draw, slide)
    for shape in slide.shapes:
        _draw_shape(draw, shape, img)
    out_path = out_dir / f"slide_{idx + 1:02d}.png"
    img.save(out_path, "PNG")
    # Half-size thumbnail
    thumb = img.resize((SLIDE_W // 2, SLIDE_H // 2), Image.LANCZOS)
    thumb.save(out_dir / f"slide_{idx + 1:02d}_thumb.png", "PNG")
    return out_path


def main(pptx_path: str, out_dir: str = "preview"):
    out = Path(out_dir)
    out.mkdir(exist_ok=True, parents=True)
    prs = Presentation(pptx_path)
    paths = []
    for i, slide in enumerate(prs.slides):
        paths.append(render_slide(slide, i, out))
    return paths


if __name__ == "__main__":
    pptx = sys.argv[1] if len(sys.argv) > 1 else "demo.pptx"
    out = sys.argv[2] if len(sys.argv) > 2 else "preview"
    for p in main(pptx, out):
        print(p)
